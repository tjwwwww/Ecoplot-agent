# -*- coding: utf-8 -*-
"""Import daily climate data into SQLite.

This module keeps the raw imported table and also maintains:
- climate_stations
- climate_daily_normalized

Rationale:
- keep source-traceable raw fields
- expose semantically normalized fields for ontology/API/agent use
"""
from __future__ import annotations

import csv
import os
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(BASE_DIR, "data", "qilian_forest.db")
ROOT_DIR = os.environ.get(
    "CLIMATE_DATA_ROOT",
    os.path.join(BASE_DIR, "data", "climate_daily_raw"),
)
RAW_TABLE_NAME = "climate_daily_observations"
NORMALIZED_TABLE_NAME = "climate_daily_normalized"
STATION_TABLE_NAME = "climate_stations"
MONTHLY_TABLE_NAME = "climate_monthly_summary"
ANNUAL_TABLE_NAME = "climate_annual_summary"
MISSING_VALUES = {"-9999", "-9999.0", "-9999.0000000", "9999.9", "999.9", "9999", "", None}

RAW_COLUMNS = [
    "station_id", "observation_date", "latitude", "longitude", "elevation_m", "station_name",
    "temp", "temp_attributes", "dewp", "dewp_attributes", "slp", "slp_attributes",
    "stp", "stp_attributes", "visib", "visib_attributes", "wdsp", "wdsp_attributes",
    "mxspd", "gust", "max_temp", "max_attributes", "min_temp", "min_attributes",
    "prcp", "prcp_attributes", "sndp", "frshtt", "source_file", "imported_at",
]


def normalize_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text in MISSING_VALUES:
        return None
    return text



def to_float(value: Any) -> Optional[float]:
    text = normalize_text(value)
    if text is None:
        return None
    try:
        return float(text)
    except ValueError:
        return None





def normalize_numeric(value: Optional[float]) -> Optional[float]:
    if value is None:
        return None
    if value in {999.9, 9999.9, -9999.0, -9999.0}:
        return None
    return value

def fahrenheit_to_celsius(value: Optional[float]) -> Optional[float]:
    value = normalize_numeric(value)
    if value is None:
        return None
    return round((value - 32.0) * 5.0 / 9.0, 4)



def knots_to_m_s(value: Optional[float]) -> Optional[float]:
    value = normalize_numeric(value)
    if value is None:
        return None
    return round(value * 0.514444, 4)



def miles_to_km(value: Optional[float]) -> Optional[float]:
    value = normalize_numeric(value)
    if value is None:
        return None
    return round(value * 1.609344, 4)



def inches_to_mm(value: Optional[float]) -> Optional[float]:
    value = normalize_numeric(value)
    if value is None:
        return None
    return round(value * 25.4, 4)



def inches_to_cm(value: Optional[float]) -> Optional[float]:
    value = normalize_numeric(value)
    if value is None:
        return None
    return round(value * 2.54, 4)



def decode_frshtt(flags: Optional[str]) -> Optional[str]:
    text = normalize_text(flags)
    if text is None:
        return None
    text = text.zfill(6)
    labels: List[str] = []
    mapping = [
        (0, "fog"),
        (1, "rain"),
        (2, "snow"),
        (3, "hail"),
        (4, "thunder"),
        (5, "tornado"),
    ]
    for idx, label in mapping:
        if idx < len(text) and text[idx] == "1":
            labels.append(label)
    return ",".join(labels) if labels else None



def build_quality_flag(row: Dict[str, Any]) -> Optional[str]:
    flags: List[str] = []
    if row.get("temp_attributes"):
        flags.append(f"TEMP_ATTR:{row['temp_attributes']}")
    if row.get("prcp_attributes"):
        flags.append(f"PRCP_ATTR:{row['prcp_attributes']}")
    if row.get("wdsp_attributes"):
        flags.append(f"WDSP_ATTR:{row['wdsp_attributes']}")
    return ";".join(flags) if flags else None



def ensure_raw_schema(conn: sqlite3.Connection) -> None:
    conn.execute(f"""
    CREATE TABLE IF NOT EXISTS {RAW_TABLE_NAME} (
        station_id TEXT,
        observation_date TEXT,
        latitude REAL,
        longitude REAL,
        elevation_m REAL,
        station_name TEXT,
        temp REAL,
        temp_attributes TEXT,
        dewp REAL,
        dewp_attributes TEXT,
        slp REAL,
        slp_attributes TEXT,
        stp REAL,
        stp_attributes TEXT,
        visib REAL,
        visib_attributes TEXT,
        wdsp REAL,
        wdsp_attributes TEXT,
        mxspd REAL,
        gust REAL,
        max_temp REAL,
        max_attributes TEXT,
        min_temp REAL,
        min_attributes TEXT,
        prcp REAL,
        prcp_attributes TEXT,
        sndp REAL,
        frshtt TEXT,
        source_file TEXT,
        imported_at TEXT,
        PRIMARY KEY (station_id, observation_date)
    )
    """)
    conn.execute(f"CREATE INDEX IF NOT EXISTS idx_{RAW_TABLE_NAME}_station_date ON {RAW_TABLE_NAME}(station_id, observation_date)")
    conn.execute(f"CREATE INDEX IF NOT EXISTS idx_{RAW_TABLE_NAME}_date ON {RAW_TABLE_NAME}(observation_date)")



def ensure_normalized_schema(conn: sqlite3.Connection) -> None:
    conn.execute(f"""
    CREATE TABLE IF NOT EXISTS {STATION_TABLE_NAME} (
        station_id TEXT PRIMARY KEY,
        station_name TEXT,
        latitude REAL,
        longitude REAL,
        elevation_m REAL,
        data_source TEXT,
        imported_at TEXT
    )
    """)
    conn.execute(f"""
    CREATE TABLE IF NOT EXISTS {NORMALIZED_TABLE_NAME} (
        station_id TEXT NOT NULL,
        observation_date TEXT NOT NULL,
        mean_temperature_c REAL,
        min_temperature_c REAL,
        max_temperature_c REAL,
        precipitation_mm REAL,
        dew_point_c REAL,
        wind_speed_m_s REAL,
        wind_gust_m_s REAL,
        max_wind_speed_m_s REAL,
        visibility_km REAL,
        sea_level_pressure_hpa REAL,
        station_pressure_hpa REAL,
        snow_depth_cm REAL,
        raw_weather_flags TEXT,
        weather_type TEXT,
        quality_flag TEXT,
        data_source TEXT,
        imported_at TEXT,
        PRIMARY KEY (station_id, observation_date)
    )
    """)
    conn.execute(f"""
    CREATE TABLE IF NOT EXISTS {MONTHLY_TABLE_NAME} (
        station_id TEXT NOT NULL,
        year INTEGER NOT NULL,
        month INTEGER NOT NULL,
        record_count INTEGER,
        valid_mean_temperature_days INTEGER,
        valid_precipitation_days INTEGER,
        mean_temperature_c REAL,
        mean_min_temperature_c REAL,
        mean_max_temperature_c REAL,
        absolute_min_temperature_c REAL,
        absolute_max_temperature_c REAL,
        total_precipitation_mm REAL,
        mean_dew_point_c REAL,
        mean_wind_speed_m_s REAL,
        max_wind_gust_m_s REAL,
        max_wind_speed_m_s REAL,
        mean_visibility_km REAL,
        mean_sea_level_pressure_hpa REAL,
        mean_station_pressure_hpa REAL,
        mean_snow_depth_cm REAL,
        rain_day_count INTEGER,
        snow_day_count INTEGER,
        thunder_day_count INTEGER,
        data_source TEXT,
        imported_at TEXT,
        PRIMARY KEY (station_id, year, month)
    )
    """)
    conn.execute(f"""
    CREATE TABLE IF NOT EXISTS {ANNUAL_TABLE_NAME} (
        station_id TEXT NOT NULL,
        year INTEGER NOT NULL,
        record_count INTEGER,
        valid_mean_temperature_days INTEGER,
        valid_precipitation_days INTEGER,
        mean_temperature_c REAL,
        mean_min_temperature_c REAL,
        mean_max_temperature_c REAL,
        absolute_min_temperature_c REAL,
        absolute_max_temperature_c REAL,
        total_precipitation_mm REAL,
        mean_dew_point_c REAL,
        mean_wind_speed_m_s REAL,
        max_wind_gust_m_s REAL,
        max_wind_speed_m_s REAL,
        mean_visibility_km REAL,
        mean_sea_level_pressure_hpa REAL,
        mean_station_pressure_hpa REAL,
        mean_snow_depth_cm REAL,
        rain_day_count INTEGER,
        snow_day_count INTEGER,
        thunder_day_count INTEGER,
        data_source TEXT,
        imported_at TEXT,
        PRIMARY KEY (station_id, year)
    )
    """)
    conn.execute(f"CREATE INDEX IF NOT EXISTS idx_{NORMALIZED_TABLE_NAME}_station_date ON {NORMALIZED_TABLE_NAME}(station_id, observation_date)")
    conn.execute(f"CREATE INDEX IF NOT EXISTS idx_{NORMALIZED_TABLE_NAME}_date ON {NORMALIZED_TABLE_NAME}(observation_date)")
    conn.execute(f"CREATE INDEX IF NOT EXISTS idx_{NORMALIZED_TABLE_NAME}_weather_type ON {NORMALIZED_TABLE_NAME}(weather_type)")
    conn.execute(f"CREATE INDEX IF NOT EXISTS idx_{MONTHLY_TABLE_NAME}_station_year_month ON {MONTHLY_TABLE_NAME}(station_id, year, month)")
    conn.execute(f"CREATE INDEX IF NOT EXISTS idx_{ANNUAL_TABLE_NAME}_station_year ON {ANNUAL_TABLE_NAME}(station_id, year)")



def ensure_schema(conn: sqlite3.Connection) -> None:
    ensure_raw_schema(conn)
    ensure_normalized_schema(conn)
    conn.commit()



def discover_station_files(root: str) -> List[Path]:
    if not os.path.isdir(root):
        return []
    station_files: List[Path] = []
    for year_dir in sorted(os.listdir(root)):
        year_path = Path(root) / year_dir
        if not year_path.is_dir():
            continue
        for candidate in sorted(year_path.iterdir()):
            if candidate.suffix.lower() not in {".csv", ".xls", ".xlsx"}:
                continue
            if candidate.name.startswith("~$"):
                continue
            station_files.append(candidate)
    return station_files



def read_csv_rows(file_path: Path) -> Iterable[Dict[str, Any]]:
    with file_path.open("r", encoding="utf-8", errors="ignore") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames is None:
            return []
        for row in reader:
            yield row



def read_excel_rows(file_path: Path) -> Iterable[Dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read Excel climate files") from exc
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(x).strip() if x is not None else "" for x in rows[0]]
    for row in rows[1:]:
        yield {headers[i]: row[i] for i in range(min(len(headers), len(row)))}



def load_station_rows(file_path: Path) -> Iterable[Dict[str, Any]]:
    if file_path.suffix.lower() == ".csv":
        yield from read_csv_rows(file_path)
    elif file_path.suffix.lower() in {".xls", ".xlsx"}:
        yield from read_excel_rows(file_path)



def normalize_station_row(raw: Dict[str, Any], file_path: Path) -> Dict[str, Any]:
    observation_date = normalize_text(raw.get("DATE") or raw.get("date") or raw.get("Date"))
    if observation_date and len(observation_date) == 8 and observation_date.isdigit():
        observation_date = f"{observation_date[:4]}-{observation_date[4:6]}-{observation_date[6:8]}"
    return {
        "station_id": normalize_text(raw.get("STATION") or raw.get("station") or raw.get("Station")),
        "observation_date": observation_date,
        "latitude": to_float(raw.get("LATITUDE") or raw.get("latitude") or raw.get("Lat")),
        "longitude": to_float(raw.get("LONGITUDE") or raw.get("longitude") or raw.get("Lon")),
        "elevation_m": to_float(raw.get("ELEVATION") or raw.get("elevation") or raw.get("Elev")),
        "station_name": normalize_text(raw.get("NAME") or raw.get("name") or raw.get("Station Name")),
        "temp": to_float(raw.get("TEMP") or raw.get("temp")),
        "temp_attributes": normalize_text(raw.get("TEMP_ATTRIBUTES") or raw.get("temp_attributes")),
        "dewp": to_float(raw.get("DEWP") or raw.get("dewp")),
        "dewp_attributes": normalize_text(raw.get("DEWP_ATTRIBUTES") or raw.get("dewp_attributes")),
        "slp": to_float(raw.get("SLP") or raw.get("slp")),
        "slp_attributes": normalize_text(raw.get("SLP_ATTRIBUTES") or raw.get("slp_attributes")),
        "stp": to_float(raw.get("STP") or raw.get("stp")),
        "stp_attributes": normalize_text(raw.get("STP_ATTRIBUTES") or raw.get("stp_attributes")),
        "visib": to_float(raw.get("VISIB") or raw.get("visib")),
        "visib_attributes": normalize_text(raw.get("VISIB_ATTRIBUTES") or raw.get("visib_attributes")),
        "wdsp": to_float(raw.get("WDSP") or raw.get("wdsp")),
        "wdsp_attributes": normalize_text(raw.get("WDSP_ATTRIBUTES") or raw.get("wdsp_attributes")),
        "mxspd": to_float(raw.get("MXSPD") or raw.get("mxspd")),
        "gust": to_float(raw.get("GUST") or raw.get("gust")),
        "max_temp": to_float(raw.get("MAX") or raw.get("max")),
        "max_attributes": normalize_text(raw.get("MAX_ATTRIBUTES") or raw.get("max_attributes")),
        "min_temp": to_float(raw.get("MIN") or raw.get("min")),
        "min_attributes": normalize_text(raw.get("MIN_ATTRIBUTES") or raw.get("min_attributes")),
        "prcp": to_float(raw.get("PRCP") or raw.get("prcp")),
        "prcp_attributes": normalize_text(raw.get("PRCP_ATTRIBUTES") or raw.get("prcp_attributes")),
        "sndp": to_float(raw.get("SNDP") or raw.get("sndp")),
        "frshtt": normalize_text(raw.get("FRSHTT") or raw.get("frshtt")),
        "source_file": file_path.name,
        "imported_at": datetime.utcnow().isoformat() + "Z",
    }



def upsert_raw_row(conn: sqlite3.Connection, row: Dict[str, Any]) -> None:
    placeholders = ", ".join("?" for _ in RAW_COLUMNS)
    conn.execute(
        f"REPLACE INTO {RAW_TABLE_NAME} ({', '.join(RAW_COLUMNS)}) VALUES ({placeholders})",
        [row.get(col) for col in RAW_COLUMNS],
    )



def upsert_station(conn: sqlite3.Connection, row: Dict[str, Any]) -> None:
    if not row.get("station_id"):
        return
    conn.execute(
        f"""
        REPLACE INTO {STATION_TABLE_NAME}
        (station_id, station_name, latitude, longitude, elevation_m, data_source, imported_at)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            row.get("station_id"),
            row.get("station_name"),
            row.get("latitude"),
            row.get("longitude"),
            row.get("elevation_m"),
            row.get("source_file"),
            row.get("imported_at"),
        ),
    )



def normalized_row_from_raw(row: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "station_id": row.get("station_id"),
        "observation_date": row.get("observation_date"),
        "mean_temperature_c": fahrenheit_to_celsius(row.get("temp")),
        "min_temperature_c": fahrenheit_to_celsius(row.get("min_temp")),
        "max_temperature_c": fahrenheit_to_celsius(row.get("max_temp")),
        "precipitation_mm": inches_to_mm(row.get("prcp")),
        "dew_point_c": fahrenheit_to_celsius(row.get("dewp")),
        "wind_speed_m_s": knots_to_m_s(row.get("wdsp")),
        "wind_gust_m_s": knots_to_m_s(row.get("gust")),
        "max_wind_speed_m_s": knots_to_m_s(row.get("mxspd")),
        "visibility_km": miles_to_km(row.get("visib")),
        "sea_level_pressure_hpa": normalize_numeric(row.get("slp")),
        "station_pressure_hpa": normalize_numeric(row.get("stp")),
        "snow_depth_cm": inches_to_cm(row.get("sndp")),
        "raw_weather_flags": row.get("frshtt"),
        "weather_type": decode_frshtt(row.get("frshtt")),
        "quality_flag": build_quality_flag(row),
        "data_source": row.get("source_file"),
        "imported_at": row.get("imported_at"),
    }



def upsert_normalized_row(conn: sqlite3.Connection, row: Dict[str, Any]) -> None:
    normalized = normalized_row_from_raw(row)
    conn.execute(
        f"""
        REPLACE INTO {NORMALIZED_TABLE_NAME} (
            station_id, observation_date, mean_temperature_c, min_temperature_c, max_temperature_c,
            precipitation_mm, dew_point_c, wind_speed_m_s, wind_gust_m_s, max_wind_speed_m_s,
            visibility_km, sea_level_pressure_hpa, station_pressure_hpa, snow_depth_cm,
            raw_weather_flags, weather_type, quality_flag, data_source, imported_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            normalized["station_id"], normalized["observation_date"], normalized["mean_temperature_c"],
            normalized["min_temperature_c"], normalized["max_temperature_c"], normalized["precipitation_mm"],
            normalized["dew_point_c"], normalized["wind_speed_m_s"], normalized["wind_gust_m_s"],
            normalized["max_wind_speed_m_s"], normalized["visibility_km"], normalized["sea_level_pressure_hpa"],
            normalized["station_pressure_hpa"], normalized["snow_depth_cm"], normalized["raw_weather_flags"],
            normalized["weather_type"], normalized["quality_flag"], normalized["data_source"], normalized["imported_at"],
        ),
    )



def rebuild_climate_aggregations(conn: sqlite3.Connection) -> None:
    ensure_normalized_schema(conn)
    conn.execute(f"DELETE FROM {MONTHLY_TABLE_NAME}")
    conn.execute(f"DELETE FROM {ANNUAL_TABLE_NAME}")
    monthly_sql = f"""
        INSERT INTO {MONTHLY_TABLE_NAME} (
            station_id, year, month, record_count, valid_mean_temperature_days, valid_precipitation_days,
            mean_temperature_c, mean_min_temperature_c, mean_max_temperature_c,
            absolute_min_temperature_c, absolute_max_temperature_c, total_precipitation_mm,
            mean_dew_point_c, mean_wind_speed_m_s, max_wind_gust_m_s, max_wind_speed_m_s,
            mean_visibility_km, mean_sea_level_pressure_hpa, mean_station_pressure_hpa, mean_snow_depth_cm,
            rain_day_count, snow_day_count, thunder_day_count, data_source, imported_at
        )
        SELECT
            station_id,
            CAST(strftime('%Y', observation_date) AS INTEGER) AS year,
            CAST(strftime('%m', observation_date) AS INTEGER) AS month,
            COUNT(*) AS record_count,
            COUNT(mean_temperature_c) AS valid_mean_temperature_days,
            COUNT(precipitation_mm) AS valid_precipitation_days,
            ROUND(AVG(mean_temperature_c), 4) AS mean_temperature_c,
            ROUND(AVG(min_temperature_c), 4) AS mean_min_temperature_c,
            ROUND(AVG(max_temperature_c), 4) AS mean_max_temperature_c,
            ROUND(MIN(min_temperature_c), 4) AS absolute_min_temperature_c,
            ROUND(MAX(max_temperature_c), 4) AS absolute_max_temperature_c,
            ROUND(SUM(COALESCE(precipitation_mm, 0.0)), 4) AS total_precipitation_mm,
            ROUND(AVG(dew_point_c), 4) AS mean_dew_point_c,
            ROUND(AVG(wind_speed_m_s), 4) AS mean_wind_speed_m_s,
            ROUND(MAX(wind_gust_m_s), 4) AS max_wind_gust_m_s,
            ROUND(MAX(max_wind_speed_m_s), 4) AS max_wind_speed_m_s,
            ROUND(AVG(visibility_km), 4) AS mean_visibility_km,
            ROUND(AVG(sea_level_pressure_hpa), 4) AS mean_sea_level_pressure_hpa,
            ROUND(AVG(station_pressure_hpa), 4) AS mean_station_pressure_hpa,
            ROUND(AVG(snow_depth_cm), 4) AS mean_snow_depth_cm,
            SUM(CASE WHEN instr(COALESCE(weather_type, ''), 'rain') > 0 THEN 1 ELSE 0 END) AS rain_day_count,
            SUM(CASE WHEN instr(COALESCE(weather_type, ''), 'snow') > 0 THEN 1 ELSE 0 END) AS snow_day_count,
            SUM(CASE WHEN instr(COALESCE(weather_type, ''), 'thunder') > 0 THEN 1 ELSE 0 END) AS thunder_day_count,
            MIN(data_source) AS data_source,
            MAX(imported_at) AS imported_at
        FROM {NORMALIZED_TABLE_NAME}
        GROUP BY station_id, CAST(strftime('%Y', observation_date) AS INTEGER), CAST(strftime('%m', observation_date) AS INTEGER)
    """
    annual_sql = f"""
        INSERT INTO {ANNUAL_TABLE_NAME} (
            station_id, year, record_count, valid_mean_temperature_days, valid_precipitation_days,
            mean_temperature_c, mean_min_temperature_c, mean_max_temperature_c,
            absolute_min_temperature_c, absolute_max_temperature_c, total_precipitation_mm,
            mean_dew_point_c, mean_wind_speed_m_s, max_wind_gust_m_s, max_wind_speed_m_s,
            mean_visibility_km, mean_sea_level_pressure_hpa, mean_station_pressure_hpa, mean_snow_depth_cm,
            rain_day_count, snow_day_count, thunder_day_count, data_source, imported_at
        )
        SELECT
            station_id,
            CAST(strftime('%Y', observation_date) AS INTEGER) AS year,
            COUNT(*) AS record_count,
            COUNT(mean_temperature_c) AS valid_mean_temperature_days,
            COUNT(precipitation_mm) AS valid_precipitation_days,
            ROUND(AVG(mean_temperature_c), 4) AS mean_temperature_c,
            ROUND(AVG(min_temperature_c), 4) AS mean_min_temperature_c,
            ROUND(AVG(max_temperature_c), 4) AS mean_max_temperature_c,
            ROUND(MIN(min_temperature_c), 4) AS absolute_min_temperature_c,
            ROUND(MAX(max_temperature_c), 4) AS absolute_max_temperature_c,
            ROUND(SUM(COALESCE(precipitation_mm, 0.0)), 4) AS total_precipitation_mm,
            ROUND(AVG(dew_point_c), 4) AS mean_dew_point_c,
            ROUND(AVG(wind_speed_m_s), 4) AS mean_wind_speed_m_s,
            ROUND(MAX(wind_gust_m_s), 4) AS max_wind_gust_m_s,
            ROUND(MAX(max_wind_speed_m_s), 4) AS max_wind_speed_m_s,
            ROUND(AVG(visibility_km), 4) AS mean_visibility_km,
            ROUND(AVG(sea_level_pressure_hpa), 4) AS mean_sea_level_pressure_hpa,
            ROUND(AVG(station_pressure_hpa), 4) AS mean_station_pressure_hpa,
            ROUND(AVG(snow_depth_cm), 4) AS mean_snow_depth_cm,
            SUM(CASE WHEN instr(COALESCE(weather_type, ''), 'rain') > 0 THEN 1 ELSE 0 END) AS rain_day_count,
            SUM(CASE WHEN instr(COALESCE(weather_type, ''), 'snow') > 0 THEN 1 ELSE 0 END) AS snow_day_count,
            SUM(CASE WHEN instr(COALESCE(weather_type, ''), 'thunder') > 0 THEN 1 ELSE 0 END) AS thunder_day_count,
            MIN(data_source) AS data_source,
            MAX(imported_at) AS imported_at
        FROM {NORMALIZED_TABLE_NAME}
        GROUP BY station_id, CAST(strftime('%Y', observation_date) AS INTEGER)
    """
    conn.execute(monthly_sql)
    conn.execute(annual_sql)
    conn.commit()


def rebuild_normalized_tables(conn: sqlite3.Connection) -> int:
    ensure_normalized_schema(conn)
    rows = conn.execute(f"SELECT {', '.join(RAW_COLUMNS)} FROM {RAW_TABLE_NAME}").fetchall()
    columns = RAW_COLUMNS
    rebuilt = 0
    for raw in rows:
        row = dict(zip(columns, raw))
        upsert_station(conn, row)
        upsert_normalized_row(conn, row)
        rebuilt += 1
    conn.commit()
    rebuild_climate_aggregations(conn)
    return rebuilt



def import_climate_daily() -> None:
    if not os.path.exists(DB_PATH):
        raise FileNotFoundError(f"Target database not found: {DB_PATH}")
    if not os.path.isdir(ROOT_DIR):
        raise FileNotFoundError(f"Climate source directory not found: {ROOT_DIR}")
    station_files = discover_station_files(ROOT_DIR)
    if not station_files:
        raise FileNotFoundError(f"No station files found: {ROOT_DIR}")
    with sqlite3.connect(DB_PATH) as conn:
        ensure_schema(conn)
        inserted = 0
        skipped = 0
        for file_path in station_files:
            print(f"Processing {file_path}")
            try:
                for raw in load_station_rows(file_path):
                    row = normalize_station_row(raw, file_path)
                    if not row["station_id"] or not row["observation_date"]:
                        skipped += 1
                        continue
                    upsert_raw_row(conn, row)
                    upsert_station(conn, row)
                    upsert_normalized_row(conn, row)
                    inserted += 1
            except Exception as exc:
                print(f"WARN: failed to import {file_path}: {exc}")
        conn.commit()
        rebuild_climate_aggregations(conn)
    print(f"Imported raw rows: {inserted}; skipped rows: {skipped}")
    print(f"Database: {DB_PATH}")
    print(f"Raw table: {RAW_TABLE_NAME}")
    print(f"Normalized table: {NORMALIZED_TABLE_NAME}")
    print(f"Station table: {STATION_TABLE_NAME}")
    print(f"Monthly summary table: {MONTHLY_TABLE_NAME}")
    print(f"Annual summary table: {ANNUAL_TABLE_NAME}")


if __name__ == "__main__":
    import_climate_daily()
